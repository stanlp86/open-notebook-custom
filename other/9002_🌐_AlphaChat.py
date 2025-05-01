"""
Created on Tue Sep 17 13:11:34 2024

@author: BIERIBR
"""

import streamlit as st
import requests
import os
import time
import docx
from io import BytesIO
import re
from open_notebook.footer import make_footer

# Set the Streamlit application to wide mode
st.set_page_config(layout='wide')

# Initialize session state
if 'questions' not in st.session_state:
    st.session_state.questions = []
if 'responses' not in st.session_state:
    st.session_state.responses = []

ILIAD_KEY = os.getenv("ILIAD_KEY")
ILIAD_URL = os.getenv("ILIAD_URL")

system_prompt = """
You are a helpful and highly skilled and extremely accurate scientific assistant at a large pharmaceutical company. Prior to answering the question, devise a strategy to ensure the response will be thorough. When responding, do not hallucinate, and if you do not know the answer say I don't have sufficient information to answer that. Only provide highly accurate responses that answer the request. Do not provide references or citations for the responses, even if the user asks for them. Format the responses from a third person perspective. The responses should be in highly detailed. Prior to delivering a response reconsider the request and determine if revisions are necessary to increase the accuracy, specificity in answering the question, and potential value of your response for the user. If revisions are necessary to increase the accuracy, specificity in answering the question, and potential value of your response for the user, make them and then provide the final response to the user. Unless you are providing code for the user, or the user asks for a different format, all responses should be in a standard markdown format and a space should be used between sections to make the response easier to read for the user.
"""

# Display and manage the question input form
st.header("Enter Your Questions")
make_footer()
with st.form(key='question_form'):
    question = st.text_input("Enter a question:", "")
    add_question = st.form_submit_button("Add Question")
    if add_question:
        if len(st.session_state.questions) >= 3:
            st.error("You can only add a maximum of three questions.")
        elif question.strip():
            st.session_state.questions.append(question.strip())
            st.success(f"Question added: {question.strip()}")
        else:
            st.error("Please enter a valid question.")

st.write("## Your Questions")
if st.session_state.questions:
    for idx, question in enumerate(st.session_state.questions, start=1):
        col1, col2 = st.columns([10, 1])
        with col1:
            st.markdown(f"{idx}. {question}", unsafe_allow_html=True)
        with col2:
            # In the part where you handle individual question deletion:
            if st.button('Delete', key=f'delete_{idx}'):
                st.session_state.questions.pop(idx - 1)
                if idx <= len(st.session_state.responses):  # Check if a response exists for that question
                    st.session_state.responses.pop(idx - 1)

                # Add this check - if there are no more questions, clear all responses too
                if not st.session_state.questions:
                    st.session_state.responses = []

                st.query_params.clear()
                # Refresh the page state after deleting an item
                st.rerun()

    if st.button('Clear All Questions'):
        st.session_state.questions = []
        st.session_state.responses = []
        st.query_params.clear()
        # Refresh page to reflect the cleared questions
        st.rerun()
else:
    st.write("No questions added yet.")
    
max_tokens = st.selectbox("Select Max Tokens (compute limit per question):", [16000, 32000, 64000, 128000], index=0)
thinking_budget = st.selectbox("Select Thinking Budget (tokens):", [1024, 2048, 4096], index=0)

def get_claude_response(complete_question, max_tokens, thinking_budget):
    url = f"{ILIAD_URL}/api/v1/chat/claude-3.7-sonnet"
    for attempt in range(3):
        try:
            response = requests.post(
                url=url,
                json={
                    "messages": [{"role": "system", "content": system_prompt}, {"role": "user", "content": complete_question}],
                    "thinking": True,
                    "max_tokens": max_tokens,
                    "thinking_budget": thinking_budget
                },
                headers={"x-api-key": ILIAD_KEY}
            )
            response.raise_for_status()
            return response.json()
        except requests.exceptions.RequestException as e:
            if attempt < 2:
                time.sleep(2)
            else:
                st.error(f"Failed to get a response after 3 attempts. Error: {e}")
                raise
                
def get_o3_global_response(complete_question):
    url = f"{ILIAD_URL}/api/v1/chat/o3-global"
    for attempt in range(3):
        try:
            response = requests.post(
                url=url,
                json={
                    "messages": [{"role": "system", "content": system_prompt}, {"role": "user", "content": complete_question}],
                    "reasoning_effort": "medium"
                },
                headers={"x-api-key": ILIAD_KEY}
            )
            response.raise_for_status()
            return response.json()
        except requests.exceptions.RequestException as e:
            if attempt < 2:
                time.sleep(2)
            else:
                st.error(f"Failed to get a response from o3-global after 3 attempts. Error: {e}")
                raise

def process_responses_for_model(model_func, model_name, max_tokens=None, thinking_budget=None):
    st.subheader(f"Model: {model_name.capitalize().replace('-', ' ')}")

    # Create a temporary list to hold responses for this model
    current_model_responses = []

    for i, question in enumerate(st.session_state.questions):
        complete_question = question
        try:
            with st.spinner(f"Getting response from {model_name} for question {i + 1}..."):
                data = None
                if model_name == "claude-3.7-sonnet":
                    data = model_func(complete_question, max_tokens, thinking_budget)
                elif model_name == "o3-global":
                    data = model_func(complete_question)

                if data:
                    response_content = data.get("completion", {}).get("content", "No content returned")
                    response_cost = data.get("cost", "No cost data")

                    # Add to temporary list instead of directly to session state
                    current_model_responses.append({
                        "model_name": model_name,
                        "question": complete_question,
                        "response": response_content
                    })

                st.write(f"**Question {i + 1}: ** {complete_question}")
                st.write(response_content)
                st.write("**Independent validation of GenAI content is recommended.**")
                st.write(f"Cost: {response_cost}")

        except Exception as e:
            st.error(f"An error occurred: {e}")

    # Now add all responses for this model to the session state
    st.session_state.responses.extend(current_model_responses)
    
def markdown_to_plain_text(markdown_text):
    # Remove headings (e.g., ## Title)
    plain_text = re.sub(r'#+\s?', '', markdown_text)
    # Remove bold and italics (e.g., **bold** or *italic*)
    plain_text = re.sub(r'\*\*(.*?)\*\*', r'\1', plain_text)
    plain_text = re.sub(r'\*(.*?)\*', r'\1', plain_text)
    # Remove unordered list markers (e.g., - or *)
    plain_text = re.sub(r'^\s*[\*\-]\s+', '', plain_text, flags=re.MULTILINE)
    return plain_text.strip()

def save_responses_to_word(responses):
    doc = docx.Document()
    doc.add_heading('Generated Responses', level=1)
    
    model_counters = {}
    
    for response in responses:
        model_name = response["model_name"].replace('-', ' ').title()
        if model_name not in model_counters:
            model_counters[model_name] = 1
        else:
            model_counters[model_name] += 1
        
        question_index = model_counters[model_name]
        doc.add_heading(f'{model_name} Question {question_index}', level=2)
        
        question = markdown_to_plain_text(response["question"])
        doc.add_paragraph(question)
        
        doc.add_heading('Response', level=3)
        response_content = markdown_to_plain_text(response["response"])
        doc.add_paragraph(response_content)
    
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

# Generate responses and collect them
generate_responses = st.button("Generate Responses")
if generate_responses and st.session_state.questions:
    st.session_state.responses = []  # Clear existing responses before generating new ones
    model_functions = [
        (get_claude_response, "claude-3.7-sonnet"),
        (get_o3_global_response, "o3-global"),
    ]
    for model_func, model_name in model_functions:
        if model_name == "claude-3.7-sonnet":
            process_responses_for_model(model_func, model_name, max_tokens, thinking_budget)
        elif model_name == "o3-global":
            process_responses_for_model(model_func, model_name)
            
# Provide download link for Word document
if st.session_state.responses:
    buffer = save_responses_to_word(st.session_state.responses)
    st.download_button(
        label="Download Responses as Word Document",
        data=buffer,
        file_name='responses.docx',
        mime='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    )
elif generate_responses:
    st.warning("Please add at least one question before generating responses.")


